#!/usr/bin/env python3
"""Génère la page « kit partenaire » d'un échange d'articles, sur le modèle de celle de Sorank.

Entrée  : kits/<partenaire>/kit.json + article.md + images/
Sortie  : kit/<id>/ — page privée (noindex) avec l'article prêt à copier, la balise title,
          la meta description, les images, et les fichiers Word, HTML, Markdown et zip.

Usage   : python3 scripts/build-kit.py kits/sedestral
Requis  : pip3 install python-docx
"""
import html
import json
import re
import shutil
import sys
import zipfile
from pathlib import Path

from docx import Document
from docx.shared import Inches, Pt

ROOT = Path(__file__).resolve().parent.parent
SITE = "https://msd-media.com"


# ── Markdown → HTML (sous-ensemble utilisé par nos articles) ────────────────────────────────
def inline(text):
    out = html.escape(text, quote=False)
    out = re.sub(r"\*\*([^*]+)\*\*", r"<strong>\1</strong>", out)
    out = re.sub(r"(?<![*\w])\*([^*]+)\*(?!\w)", r"<em>\1</em>", out)
    out = re.sub(r"\[([^\]]+)\]\(([^)\s]+)\)", r'<a href="\2">\1</a>', out)
    return out


def slugify(text):
    import unicodedata
    t = unicodedata.normalize("NFD", text.lower())
    t = "".join(c for c in t if unicodedata.category(c) != "Mn")
    return re.sub(r"[^a-z0-9]+", "-", t).strip("-")


def md_to_blocks(md):
    """Découpe le Markdown en blocs typés, réutilisés pour le HTML et le Word."""
    lines = md.replace("\r\n", "\n").split("\n")
    blocks, i = [], 0
    while i < len(lines):
        line = lines[i]
        if not line.strip():
            i += 1
            continue
        if m := re.match(r"^(#{1,4})\s+(.+)$", line):
            blocks.append(("h", len(m.group(1)), m.group(2).strip()))
            i += 1
            continue
        if m := re.match(r'^!\[([^\]]*)\]\(\s*([^\s)]+)(?:\s+"([^"]*)")?\s*\)$', line.strip()):
            blocks.append(("img", m.group(1), m.group(2), m.group(3) or ""))
            i += 1
            continue
        if re.match(r"^\|.+\|\s*$", line) and i + 1 < len(lines) and re.match(r"^\|?\s*[-:]+", lines[i + 1]):
            rows = []
            while i < len(lines) and re.match(r"^\|.+\|\s*$", lines[i]):
                if not re.match(r"^\|?\s*[-:]+\s*(\|\s*[-:]+\s*)+\|?$", lines[i]):
                    rows.append([c.strip() for c in lines[i].strip().strip("|").split("|")])
                i += 1
            blocks.append(("table", rows))
            continue
        if re.match(r"^[-*+]\s+", line):
            items = []
            while i < len(lines) and re.match(r"^[-*+]\s+", lines[i]):
                items.append(re.sub(r"^[-*+]\s+", "", lines[i]))
                i += 1
            blocks.append(("ul", items))
            continue
        if re.match(r"^\d+\.\s+", line):
            items = []
            while i < len(lines) and re.match(r"^\d+\.\s+", lines[i]):
                items.append(re.sub(r"^\d+\.\s+", "", lines[i]))
                i += 1
            blocks.append(("ol", items))
            continue
        para = [line.strip()]
        i += 1
        while i < len(lines) and lines[i].strip() and not re.match(r"^(#{1,4}\s|[-*+]\s|\d+\.\s|\||!\[)", lines[i]):
            para.append(lines[i].strip())
            i += 1
        blocks.append(("p", " ".join(para)))
    return blocks


def blocks_to_html(blocks, img_prefix=""):
    parts = []
    for b in blocks:
        kind = b[0]
        if kind == "h":
            parts.append(f'<h{b[1]} id="{slugify(b[2])}">{inline(b[2])}</h{b[1]}>')
        elif kind == "img":
            src = img_prefix + b[2] if not b[2].startswith("http") else b[2]
            cap = f"<figcaption>{html.escape(b[3])}</figcaption>" if b[3] else ""
            parts.append(f'<figure><img src="{html.escape(src)}" alt="{html.escape(b[1])}" loading="lazy">{cap}</figure>')
        elif kind == "table":
            head, *rows = b[1]
            th = "".join(f"<th>{inline(c)}</th>" for c in head)
            trs = "".join("<tr>" + "".join(f"<td>{inline(c)}</td>" for c in r) + "</tr>" for r in rows)
            parts.append(f"<table><thead><tr>{th}</tr></thead><tbody>{trs}</tbody></table>")
        elif kind in ("ul", "ol"):
            lis = "".join(f"<li>{inline(x)}</li>" for x in b[1])
            parts.append(f"<{kind}>{lis}</{kind}>")
        else:
            parts.append(f"<p>{inline(b[1])}</p>")
    return "\n".join(parts)


# ── Word ────────────────────────────────────────────────────────────────────────────────────
def add_runs(par, text):
    """Ajoute un texte Markdown (gras, italique, liens en clair) à un paragraphe Word."""
    text = re.sub(r"\[([^\]]+)\]\(([^)\s]+)\)", r"\1 (\2)", text)
    for tok in re.split(r"(\*\*[^*]+\*\*|(?<![*\w])\*[^*]+\*(?!\w))", text):
        if not tok:
            continue
        if tok.startswith("**"):
            par.add_run(tok[2:-2]).bold = True
        elif tok.startswith("*"):
            par.add_run(tok[1:-1]).italic = True
        else:
            par.add_run(tok)


def blocks_to_docx(blocks, kit_dir, out):
    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(11)
    for b in blocks:
        kind = b[0]
        if kind == "h":
            doc.add_heading(re.sub(r"\*", "", b[2]), level=min(b[1], 4) if b[1] > 1 else 0)
        elif kind == "img":
            path = kit_dir / b[2]
            if path.exists():
                doc.add_picture(str(path), width=Inches(6.2))
            if b[3]:
                cap = doc.add_paragraph()
                cap.add_run(b[3]).italic = True
            doc.add_paragraph(f"Texte alternatif : {b[1]}").runs[0].font.size = Pt(9)
        elif kind == "table":
            rows = b[1]
            t = doc.add_table(rows=len(rows), cols=len(rows[0]))
            t.style = "Table Grid"
            for r, row in enumerate(rows):
                for c, cell in enumerate(row):
                    t.cell(r, c).text = ""
                    add_runs(t.cell(r, c).paragraphs[0], cell)
                    if r == 0:
                        for run in t.cell(r, c).paragraphs[0].runs:
                            run.bold = True
            doc.add_paragraph()
        elif kind in ("ul", "ol"):
            for x in b[1]:
                add_runs(doc.add_paragraph(style="List Bullet" if kind == "ul" else "List Number"), x)
        else:
            add_runs(doc.add_paragraph(), b[1])
    doc.save(out)


# ── Page kit ────────────────────────────────────────────────────────────────────────────────
def copy_field(label, value, hint=""):
    count = f'<span class="kit-count">{len(value)} caractères</span>' if hint == "count" else ""
    return f'''<div class="kit-field">
          <div class="kit-field__head"><span class="kit-field__label">{label}</span>{count}</div>
          <div class="kit-field__row"><code class="kit-field__value">{html.escape(value)}</code>
          <button class="kit-btn kit-btn--ghost" type="button" data-copy-text="{html.escape(value, quote=True)}">Copier</button></div>
        </div>'''


def build(kit_dir):
    kit_dir = Path(kit_dir).resolve()
    cfg = json.loads((kit_dir / "kit.json").read_text(encoding="utf-8"))
    md = (kit_dir / cfg["article"]["file"]).read_text(encoding="utf-8")
    blocks = md_to_blocks(md)
    h1 = next(b[2] for b in blocks if b[0] == "h" and b[1] == 1)
    body_blocks = [b for b in blocks if not (b[0] == "h" and b[1] == 1)]

    out = ROOT / "kit" / cfg["id"]
    if out.exists():
        shutil.rmtree(out)
    (out / "images").mkdir(parents=True)
    for img in sorted((kit_dir / "images").glob("*")):
        shutil.copy(img, out / "images" / img.name)

    # Fichiers à télécharger
    article_html = blocks_to_html(blocks, "images/").replace(' loading="lazy"', "")
    (out / "article.md").write_text(md, encoding="utf-8")
    (out / "article.html").write_text(
        f'<!doctype html>\n<html lang="fr">\n<head>\n<meta charset="utf-8">\n<title>{html.escape(cfg["article"]["title"])}</title>\n'
        f'<meta name="description" content="{html.escape(cfg["article"]["meta"])}">\n</head>\n<body>\n{article_html}\n</body>\n</html>\n',
        encoding="utf-8")
    blocks_to_docx(blocks, kit_dir, out / "article.docx")
    images = [b for b in body_blocks if b[0] == "img"]
    headings_before = {}
    last_h = ""
    for b in body_blocks:
        if b[0] == "h":
            last_h = b[2]
        elif b[0] == "img":
            headings_before[b[2]] = last_h
    readme = [f"Article pour {cfg['site']}, rédigé pour vous et publié nulle part ailleurs.", "",
              f"Titre (H1) : {h1}", f"Balise title : {cfg['article']['title']}",
              f"Meta description : {cfg['article']['meta']}", f"Adresse suggérée : {cfg['article']['slug']}", "", "Images :"]
    for n, b in enumerate(images, 1):
        readme.append(f"{n}. {Path(b[2]).name} : sous le titre « {headings_before[b[2]]} ». Texte alternatif : {b[1]}")
    readme += ["", "Deux points seulement : " + " ".join(r[0].lower() + r[1:] for r in cfg["rules"])]
    (out / "LISEZ-MOI.txt").write_text("\n".join(readme) + "\n", encoding="utf-8")
    zip_name = f"article-{slugify(cfg['partner'])}.zip"
    with zipfile.ZipFile(out / zip_name, "w", zipfile.ZIP_DEFLATED) as z:
        for f in ["article.html", "article.md", "article.docx", "LISEZ-MOI.txt"]:
            z.write(out / f, f)
        for img in sorted((out / "images").glob("*")):
            z.write(img, f"images/{img.name}")

    # Page
    footer = re.search(r'<footer class="site-footer">[\s\S]*?</footer>', (ROOT / "index.html").read_text(encoding="utf-8")).group(0)
    css_links = "\n".join(
        f'  <link rel="stylesheet" href="../../assets/css/{n}">'
        for n in dict.fromkeys(re.findall(r'href="assets/css/((?:style|responsive|fontawesome-subset)\.css\?v=[0-9a-f]+)"', (ROOT / "index.html").read_text(encoding="utf-8"))))
    listing = cfg["listing"]
    images_html = "\n".join(f'''        <article class="kit-image">
          <img src="images/{Path(b[2]).name}" alt="{html.escape(b[1])}" loading="lazy">
          <div class="kit-image__body">
            <p class="kit-image__meta">Image {n} · sous le titre « {html.escape(headings_before[b[2]])} »</p>
            <p class="kit-image__file">{Path(b[2]).name}</p>
            {copy_field("Texte alternatif", b[1])}
            <a class="kit-btn kit-btn--ghost" href="images/{Path(b[2]).name}" download>Télécharger l'image</a>
          </div>
        </article>''' for n, b in enumerate(images, 1))
    message = "\n".join(f"<p>{html.escape(p)}</p>" for p in cfg["message"])
    html_source = html.escape(article_html)
    page = f'''<!DOCTYPE html>
<html lang="fr">
<head>
  <meta charset="UTF-8" />
  <meta name="viewport" content="width=device-width, initial-scale=1.0" />
  <title>Votre article prêt à publier · {html.escape(cfg["partner"])}</title>
  <meta name="robots" content="noindex, nofollow, noarchive, nosnippet, noimageindex" />
  <link rel="icon" type="image/png" href="https://msd-media.com/assets/img/favicon-96x96.png" sizes="96x96" />
  <link rel="icon" type="image/svg+xml" href="https://msd-media.com/assets/img/favicon.svg" />
  <link rel="preconnect" href="https://fonts.googleapis.com">
  <link rel="preconnect" href="https://fonts.gstatic.com" crossorigin>
  <link href="https://fonts.googleapis.com/css2?family=Instrument+Serif:ital@0;1&family=Inter:ital,opsz,wght@0,14..32,100..900;1,14..32,100..900&display=swap" rel="stylesheet">
{css_links}
  <link rel="stylesheet" href="../kit.css">
</head>
<body class="home-page kit-page" data-kit="{cfg["id"]}">
  <header class="kit-top">
    <a href="https://msd-media.com/" aria-label="MSD Media"><img src="https://msd-media.com/assets/img/logo-black.webp" alt="MSD Media" width="372" height="151" class="kit-top__logo"></a>
    <span class="kit-top__tag">Kit partenaire</span>
  </header>

  <main class="kit">
    <section class="kit-hero">
      <p class="kit-tag">Article prêt à publier</p>
      <h1>Votre article pour {html.escape(cfg["site"])} est prêt, {html.escape(cfg["contact"])}</h1>
      <p class="kit-lead">Un article rédigé pour votre blog et publié nulle part ailleurs. Tout est prêt : il ne reste qu'à le copier.</p>
      <div class="kit-actions">
        <button class="kit-btn kit-btn--primary" type="button" data-copy-article>Copier l'article</button>
        <a class="kit-btn kit-btn--ghost" href="{zip_name}" download>Télécharger le kit (.zip)</a>
      </div>
    </section>

    <section class="kit-card kit-message">
      <img src="https://msd-media.com/assets/img/maxens-soldan.webp" alt="Maxens Soldan" class="kit-message__photo" width="96" height="96">
      <div>{message}
        <p class="kit-message__sign">Maxens Soldan, fondateur de MSD Media</p>
      </div>
    </section>

    <section class="kit-card">
      <h2>Où en est votre article</h2>
      <ol class="kit-steps">
        <li class="is-done"><span>Article rédigé pour {html.escape(cfg["site"])}</span></li>
        <li class="is-done"><span>{html.escape(cfg["partner"])} en n°{listing["rank"]} de notre comparatif</span></li>
        <li data-step-published><span>Article publié sur votre blog</span></li>
      </ol>
    </section>

    <section class="kit-card">
      <h2>Notre classement</h2>
      <a class="kit-listing" href="{listing["url"]}" target="_blank" rel="noopener">
        <img src="{listing["image"]}" alt="" loading="lazy">
        <div>
          <p class="kit-listing__label">{html.escape(listing["label"])}</p>
          <p class="kit-listing__title"><span class="kit-rank">n°{listing["rank"]}</span> {html.escape(listing["heading"])}</p>
          <p>{html.escape(listing["excerpt"])}</p>
          <p class="kit-link">Voir le comparatif →</p>
        </div>
      </a>
    </section>

    <section class="kit-card">
      <h2>Les réglages SEO</h2>
      <p class="kit-muted">À coller dans les champs de l'article.</p>
      {copy_field("Titre (H1)", h1)}
      {copy_field("Balise title", cfg["article"]["title"], "count")}
      {copy_field("Meta description", cfg["article"]["meta"], "count")}
      {copy_field("Adresse suggérée", cfg["article"]["slug"])}
    </section>

    <section class="kit-card">
      <h2>Les images</h2>
      <div class="kit-images">
{images_html}
      </div>
    </section>

    <section class="kit-card">
      <h2>Les fichiers</h2>
      <div class="kit-downloads">
        <a class="kit-btn kit-btn--ghost" href="article.docx" download>Word (.docx)</a>
        <a class="kit-btn kit-btn--ghost" href="article.html" download>HTML</a>
        <a class="kit-btn kit-btn--ghost" href="article.md" download>Markdown</a>
        <a class="kit-btn kit-btn--primary" href="{zip_name}" download>Tout le kit (.zip)</a>
      </div>
    </section>

    <section class="kit-card">
      <div class="kit-article-head">
        <h2>L'article</h2>
        <div class="kit-actions">
          <button class="kit-btn kit-btn--primary" type="button" data-copy-article>Copier l'article</button>
          <button class="kit-btn kit-btn--ghost" type="button" data-copy-html>Copier le HTML</button>
        </div>
      </div>
      <article class="kit-article" id="kit-article">
        <h1>{inline(h1)}</h1>
        {blocks_to_html(body_blocks, "images/")}
      </article>
      <textarea class="kit-hidden" id="kit-html" readonly>{html_source}</textarea>
    </section>

    <section class="kit-card" id="publie">
      <h2>Votre article est en ligne ?</h2>
      <p class="kit-muted">Collez son adresse : nous vérifions le lien et vous confirmons la mise en avant de {html.escape(cfg["partner"])}.</p>
      <form class="kit-form" data-kit-form="publish">
        <input type="url" name="url" placeholder="https://{html.escape(cfg["site"])}/blog/{html.escape(cfg["article"]["slug"])}" required>
        <button class="kit-btn kit-btn--primary" type="submit">Envoyer l'adresse</button>
      </form>
      <p class="kit-form__status" data-status="publish" role="status"></p>
    </section>

    <section class="kit-card">
      <h2>Votre fiche dans notre classement</h2>
      <p class="kit-muted">Vous voulez ajuster la description de {html.escape(cfg["partner"])} ? Proposez votre texte, nous l'intégrons.</p>
      <form class="kit-form kit-form--stack" data-kit-form="fiche">
        <textarea name="text" rows="4" placeholder="Votre description en deux ou trois phrases" required></textarea>
        <button class="kit-btn kit-btn--ghost" type="submit">Envoyer la description</button>
      </form>
      <p class="kit-form__status" data-status="fiche" role="status"></p>
    </section>
  </main>

  {footer}
  <script src="../kit.js" defer></script>
</body>
</html>
'''
    (out / "index.html").write_text(page, encoding="utf-8")
    print(f"✅ Kit {cfg['partner']} : {SITE}/kit/{cfg['id']}/  ({len(images)} images, {zip_name})")


if __name__ == "__main__":
    build(sys.argv[1] if len(sys.argv) > 1 else "kits/sedestral")
